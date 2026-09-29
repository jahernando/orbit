"""views/ledger_export.py — paquete para compartir el ledger (ADR-053, F4).

`orbit ledger <proyecto> --export <dir>` genera en *dir*:

    ledger.pdf        la verdad del usuario, para leer (colaboradores)
    ledger.xlsx       copia editable: resumen, operaciones y movimientos
    justificantes/    SOLO los ficheros enlazados por algún movimiento

Reglas:

* **Antes, `ledger --check`**: con errores no se exporta (un ledger con un
  justificante que falta no se comparte). Los avisos se cuentan y se sigue.
* **Nunca** se copia nada de `cloud/logs/` que no enlace un movimiento: allí
  hay documentos sensibles que no son del ledger.
* **Idempotente**: regenerar sobrescribe `ledger.*` y sincroniza
  `justificantes/` (añade lo nuevo, quita lo que ya no se enlaza). Solo se
  toca lo que orbit escribe: `ledger.pdf`, `ledger.xlsx` y `justificantes/`.
* Enlaces **relativos** a `justificantes/`: funcionan con la carpeta
  descargada o sincronizada.

La publicación (OneDrive, etc.) no es cosa de orbit: la hace otra herramienta
sobre *dir*.

Dependencias opcionales: `reportlab` (PDF) y `openpyxl` (xlsx) — extra
`ledger` de `pyproject.toml`.
"""

import shutil
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from core.ledger import (
    CARRY_TAG, DIRECT, INCOME_TAG, Movement, Operation, Summary,
    build_operations, currency_symbol, format_amount,
    project_partida, read_movements, summarize,
)
from views.ledger_check import ERROR, WARNING, _resolve_link, check_ledger

JUST_DIR = "justificantes"
PDF_FILE = "ledger.pdf"
XLSX_FILE = "ledger.xlsx"

_STATE_LABEL = {DIRECT: "gasto directo"}
_FUNDS = {INCOME_TAG: "Ingreso", CARRY_TAG: "Arrastre"}
_TAG_LABEL = {"gasto": "Gasto", "ingreso": "Ingreso", "arrastre": "Arrastre", "pedido": "Pedido"}


# ── Justificantes ────────────────────────────────────────────────────────────

def plan_attachments(project_dir: Path, movements: List[Movement]
                     ) -> Dict[str, Tuple[Path, str]]:
    """`{link del movimiento: (fichero origen, nombre en justificantes/)}`.

    Solo enlaces locales que existen (las URL se dejan como están). Dos
    ficheros distintos con el mismo nombre no se pisan: el segundo recibe un
    sufijo.
    """
    plan: Dict[str, Tuple[Path, str]] = {}
    taken: Dict[str, Path] = {}
    for m in movements:
        if not m.link or m.link in plan:
            continue
        src = _resolve_link(project_dir, m.link)
        if src is None or not src.is_file():
            continue
        name = src.name
        n = 2
        while name in taken and taken[name].resolve() != src.resolve():
            name = f"{src.stem}_{n}{src.suffix}"
            n += 1
        taken[name] = src
        plan[m.link] = (src, name)
    return plan


def sync_attachments(dest: Path, plan: Dict[str, Tuple[Path, str]]) -> Tuple[int, int]:
    """Deja `dest/justificantes/` igual a *plan*. Devuelve (copiados, quitados).

    Copia solo lo que cambia (tamaño o fecha) para que la carpeta compartida
    no se resincronice entera en cada export.
    """
    just = dest / JUST_DIR
    just.mkdir(parents=True, exist_ok=True)
    wanted = {name: src for src, name in plan.values()}
    copied = removed = 0
    for name, src in wanted.items():
        target = just / name
        st = src.stat()
        if (target.exists() and target.stat().st_size == st.st_size
                and int(target.stat().st_mtime) == int(st.st_mtime)):
            continue
        shutil.copy2(src, target)
        copied += 1
    for f in just.iterdir():
        if f.is_file() and f.name not in wanted:
            f.unlink()
            removed += 1
    return copied, removed


# ── Datos comunes a PDF y xlsx ───────────────────────────────────────────────

def _doc_link(m: Movement, plan) -> Optional[str]:
    if not m.link:
        return None
    if m.link in plan:
        return f"{JUST_DIR}/{plan[m.link][1]}"
    return m.link if "://" in m.link else None


def _summary_rows(s: Summary) -> List[Tuple[str, Decimal]]:
    rows = []
    if s.carried:
        rows.append(("Saldo arrastrado", s.carried))
    rows += [("Dotación", s.income), ("Gastado", s.spent),
             ("Comprometido (pedidos abiertos)", s.committed),
             ("Disponible", s.available)]
    return rows


# ── PDF ──────────────────────────────────────────────────────────────────────

def write_pdf(path: Path, title: str, partida: Optional[str], summary: Summary,
              operations: List[Operation], movements: List[Movement], plan,
              generated: datetime, status=None) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    symbol = currency_symbol()
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    right = cell.clone("right", alignment=2)
    head = cell.clone("head", fontName="Helvetica-Bold")

    def p(text, style=cell):
        return Paragraph(text, style)

    def link(text, m):
        url = _doc_link(m, plan)
        text = escape(text)
        return (f'<link href="{escape(url)}" color="blue"><u>{text}</u></link>'
                if url else text)

    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ])

    story = [Paragraph(escape(title), styles["Title"])]
    meta = f"Generado el {generated:%Y-%m-%d %H:%M}"
    if partida:
        meta = f"Partida <b>{escape(partida)}</b> · " + meta
    story += [Paragraph(meta, styles["Normal"]), Spacer(1, 4 * mm)]

    rows = [[p("Resumen", head), p(symbol, head)]]
    for label, value in _summary_rows(summary):
        bold = label == "Disponible"
        rows.append([p(f"<b>{label}</b>" if bold else label),
                     p(f"<b>{format_amount(value)}</b>" if bold
                       else format_amount(value), right)])
    t = Table(rows, colWidths=[70 * mm, 40 * mm], hAlign="LEFT")
    t.setStyle(grid)
    story += [t, Spacer(1, 5 * mm), Paragraph("Operaciones", styles["Heading2"])]

    from views.ledger import usc_cell, usc_extra_rows
    usc = status is not None and status.error is None
    if usc:
        when = status.as_of.isoformat() if status.as_of else "—"
        story.insert(2, Paragraph(
            f"Conciliado con la USC el {when}. Columna USC: ok = casa (con el "
            "nº de la USC) · !↑ = en la USC y no aquí · !↓ = aquí y no en la "
            "USC · ! = no encaja.", styles["Normal"]))
    headers = ["Fecha", "Aut.", "Factura", "Concepto y justificantes",
               "Beneficiario", "Comprometido", "Gastado", "Estado"]
    if usc:
        headers.append("USC")
    rows = [[p(h, head) for h in headers]]
    dated = []
    for op in operations:
        docs = " · ".join(link(m.concept, m) for m in op.entries)
        committed = format_amount(op.committed) if op.committed else "—"
        row = [
            p(op.date.isoformat()), p(escape(op.op_id or "—")),
            p(escape(", ".join(op.invoice_ids) or "—")), p(docs),
            p(escape(op.payee or "—")), p(committed, right),
            p(format_amount(op.spent) if op.spent else "—", right),
            p(_STATE_LABEL.get(op.state, op.state)),
        ]
        if usc:
            row.append(p(escape(usc_cell(op, status))))
        dated.append((op.date, row))
    if usc:
        extra, _funds, mark = usc_extra_rows(status)
        for d, aut, inv, concept, payee, committed, spent in extra:
            dated.append((d or date.min, [
                p(d.isoformat() if d else "—"), p(escape(aut or "—")),
                p(escape(inv or "—")), p(escape(concept)), p(escape(payee or "—")),
                p(format_amount(committed) if committed else "—", right),
                p(format_amount(spent) if spent else "—", right),
                p("—"), p(mark)]))
    dated.sort(key=lambda r: r[0])
    rows += [r for _d, r in dated]
    widths = ([18, 28, 24, 78, 38, 24, 22, 20, 21] if usc
              else [18, 28, 24, 90, 40, 26, 24, 23])
    t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
    t.setStyle(grid)
    story += [t]

    funds = [m for m in movements if m.tag in _FUNDS and not m.is_cut]
    if funds:
        story += [Spacer(1, 5 * mm), Paragraph("Dotación", styles["Heading2"])]
        rows = [[p(h, head) for h in ("Fecha", "Tipo", "Concepto", "Origen",
                                      "Importe")]]
        for m in funds:
            rows.append([p(m.date.isoformat()), p(_FUNDS[m.tag]),
                         p(link(m.concept, m)), p(escape(m.payee or "—")),
                         p(format_amount(m.amount, plus=True), right)])
        t = Table(rows, colWidths=[w * mm for w in (20, 22, 110, 50, 30)],
                  hAlign="LEFT", repeatRows=1)
        t.setStyle(grid)
        story += [t]

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm,
                               f"{title} · página {doc.page}")

    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), title=title,
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ── xlsx ─────────────────────────────────────────────────────────────────────

def write_xlsx(path: Path, title: str, partida: Optional[str], summary: Summary,
               operations: List[Operation], movements: List[Movement], plan,
               generated: datetime, status=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    money = '#,##0.00'
    bold = Font(bold=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Partida", partida or "—"])
    ws.append(["Generado", generated.strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    for label, value in _summary_rows(summary):
        ws.append([label, value])
        ws.cell(ws.max_row, 2).number_format = money
        if label == "Disponible":
            ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    from views.ledger import usc_cell, usc_extra_rows
    usc = status is not None and status.error is None
    if usc:
        ws.append(["Conciliado con la USC",
                   status.as_of.isoformat() if status.as_of else "—"])

    ws = wb.create_sheet("Operaciones")
    ws.append(["Fecha", "Aut.", "Factura", "Concepto", "Beneficiario",
               "Comprometido", "Gastado", "Pendiente", "Estado"]
              + (["USC"] if usc else []))
    for op in operations:
        ws.append([op.date, op.op_id or "", ", ".join(op.invoice_ids),
                   op.concept, op.payee or "", op.committed or None,
                   op.spent or None, op.pending or None,
                   _STATE_LABEL.get(op.state, op.state)]
                  + ([usc_cell(op, status)] if usc else []))
        row = ws.max_row
        ws.cell(row, 1).number_format = "yyyy-mm-dd"
        for c in (6, 7, 8):
            ws.cell(row, c).number_format = money
        url = _doc_link(op.entries[0], plan)
        if url:
            ws.cell(row, 4).hyperlink = url
            ws.cell(row, 4).style = "Hyperlink"
    if usc:
        extra, _funds, mark = usc_extra_rows(status)
        for d, aut, inv, concept, payee, committed, spent in extra:
            ws.append([d, aut or "", inv or "", concept, payee or "",
                       committed, spent, None, "—", mark])
            ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
            for c in (6, 7):
                ws.cell(ws.max_row, c).number_format = money

    ws = wb.create_sheet("Movimientos")
    ws.append(["Fecha", "Tipo", "Id", "Pedido", "Concepto", "Beneficiario",
               "Importe", "Justificante"])
    for m in movements:
        if m.is_cut:
            continue
        url = _doc_link(m, plan)
        ws.append([m.date, _TAG_LABEL.get(m.tag, m.tag), m.op_id or "",
                   m.ref or "", m.concept, m.payee or "",
                   m.amount if m.amount else None,
                   Path(url).name if url else ""])
        row = ws.max_row
        ws.cell(row, 1).number_format = "yyyy-mm-dd"
        ws.cell(row, 7).number_format = money
        if url:
            ws.cell(row, 8).hyperlink = url
            ws.cell(row, 8).style = "Hyperlink"

    for sheet in (wb["Operaciones"], wb["Movimientos"]):
        for c in sheet[1]:
            c.font = bold
        sheet.freeze_panes = "A2"
        for col, width in zip("ABCDEFGHIJK", (11, 16, 16, 44, 26, 14, 14, 14, 14, 16, 22)):
            sheet.column_dimensions[col].width = width
    wb.save(path)


# ── Comando ──────────────────────────────────────────────────────────────────

def export_ledger(project_dir: Path, dest: Path,
                  generated: Optional[datetime] = None) -> dict:
    """Escribe el paquete en *dest*. Lanza `ValueError` si hay errores.

    Devuelve `{"warnings": n, "copied": n, "removed": n, "attachments": n}`.
    """
    findings = check_ledger(project_dir)
    errors = [f for f in findings if f.level == ERROR]
    if errors:
        raise ValueError("el ledger tiene errores; no se exporta:\n"
                         + "\n".join(f.render() for f in errors))
    movements, _ = read_movements(project_dir)
    if not movements:
        raise ValueError("el proyecto no tiene movimientos")
    operations, _ = build_operations(movements)
    summary = summarize(movements, operations)
    from core.log import _base_name
    partida = project_partida(project_dir)
    # Sin el emoji del tipo: las fuentes base del PDF no lo tienen.
    title = f"Ledger — {_base_name(project_dir)}"
    generated = generated or datetime.now()

    from views.ledger import load_usc
    status = load_usc(project_dir, movements)
    dest.mkdir(parents=True, exist_ok=True)
    plan = plan_attachments(project_dir, movements)
    write_pdf(dest / PDF_FILE, title, partida, summary, operations, movements,
              plan, generated, status)
    write_xlsx(dest / XLSX_FILE, title, partida, summary, operations,
               movements, plan, generated, status)
    copied, removed = sync_attachments(dest, plan)
    return {"warnings": sum(1 for f in findings if f.level == WARNING),
            "copied": copied, "removed": removed, "attachments": len(plan)}


def run_ledger_export(project: str, dest: str) -> int:
    from core.log import find_project

    project_dir = find_project(project)
    if not project_dir:
        return 1
    missing = [m for m in ("reportlab", "openpyxl") if not _importable(m)]
    if missing:
        print(f"⚠️  Falta {' y '.join(missing)}: pip install {' '.join(missing)}")
        return 1
    target = Path(dest).expanduser()
    try:
        stats = export_ledger(project_dir, target)
    except ValueError as exc:
        print(f"⚠️  {exc}")
        return 1
    print(f"💶 Ledger exportado — {project_dir.name} → {target}")
    print(f"  ✓ {PDF_FILE} · {XLSX_FILE} · {JUST_DIR}/ "
          f"({stats['attachments']} justificante"
          f"{'s' if stats['attachments'] != 1 else ''}: "
          f"{stats['copied']} copiado{'s' if stats['copied'] != 1 else ''}, "
          f"{stats['removed']} quitado{'s' if stats['removed'] != 1 else ''})")
    if stats["warnings"]:
        print(f"  ⚠️  {stats['warnings']} aviso"
              f"{'s' if stats['warnings'] != 1 else ''}: "
              f"orbit ledger {project} --check")
    return 0


def _importable(name: str) -> bool:
    import importlib.util
    return importlib.util.find_spec(name) is not None
